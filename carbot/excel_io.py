"""Чтение присланного Excel и сборка итогового файла."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import Settings
from .models import Car, Evaluation

# Возможные названия колонок (в нижнем регистре, ищем вхождение)
ALIASES = {
    "brand": ["марка", "бренд", "производитель", "make", "brand"],
    "model": ["модель", "model"],
    "title": ["наименование", "автомобиль", "название", "транспортное средство", "car", "name"],
    "year": ["год выпуска", "год", "year"],
    "mileage": ["пробег", "mileage", "км"],
    "price": ["цена", "стоимость", "срс", "price", "сумма"],
    "vin": ["vin", "вин", "идентификационный"],
    "vat": ["ндс", "vat", "налог"],
}

MULTIWORD_BRANDS = [
    "land rover", "li auto", "great wall", "alfa romeo", "aston martin", "lynk & co", "rolls royce",
    "mercedes benz", "range rover", "gac motor", "jac motors",
]

VIN_RE = re.compile(r"\b[A-HJ-NPR-Z0-9]{17}\b")


def parse_number(v) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("\xa0", " ").replace(" ", "").replace(",", ".")
    m = re.search(r"-?\d+(\.\d+)?", s)
    return float(m.group()) if m else None


def parse_vat(v, default: bool) -> bool:
    if v is None or str(v).strip() == "":
        return default
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    if any(w in s for w in ("без", "нет", "не обл", "no", "false")) or s in ("0", "-"):
        return False
    if any(w in s for w in ("полн", "да", "с ндс", "вкл", "yes", "true", "20", "+")) or s == "1":
        return True
    return default


def _match(header: str, key: str) -> bool:
    h = header.lower().strip()
    return any(a == h or a in h for a in ALIASES[key])


def _detect_header(ws) -> tuple[int, dict[str, int]]:
    """Ищем в первых 20 строках строку-заголовок, где есть хотя бы цена и марка/наименование."""
    for r in range(1, min(ws.max_row, 20) + 1):
        cols: dict[str, int] = {}
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            if val is None:
                continue
            h = str(val)
            low = h.lower()
            if _match(h, "brand") and _match(h, "model") and "title" not in cols:
                cols["title"] = c  # «Марка/модель» в одной колонке
                continue
            if "сумма ндс" in low or "в т.ч" in low or low.startswith("ндс,"):
                continue  # сумма налога — не цена и не признак
            if _match(h, "price") and "без ндс" in low and "price" in cols:
                continue  # есть цена с НДС — цену без НДС не берём
            if _match(h, "price") and "с ндс" in low:
                cols["price"] = c  # цена с НДС приоритетнее
                continue
            # порядок важен: «год выпуска» не должен стать моделью, «цена с НДС» — признаком НДС и т. п.
            for key in ("vin", "year", "mileage", "price", "vat", "brand", "model", "title"):
                if key not in cols and _match(h, key):
                    cols[key] = c
                    break
        if "price" in cols and ("brand" in cols or "title" in cols or "model" in cols):
            return r, cols
    raise ValueError(
        "Не нашёл строку с заголовками. Нужны хотя бы колонки «Цена» и «Марка»/«Модель»/«Наименование»."
    )


def split_title(title: str) -> tuple[str, str]:
    t = re.sub(r"\s+", " ", title.replace("-", " ")).strip()
    low = t.lower()
    for b in MULTIWORD_BRANDS:
        if low.startswith(b + " ") or low == b:
            n = len(b.split())
            parts = title.split()
            return " ".join(parts[:n]), " ".join(parts[n:])
    parts = title.split(maxsplit=1)
    return (parts[0], parts[1] if len(parts) > 1 else "")


def read_cars(path: str | Path, s: Settings) -> list[Car]:
    wb = load_workbook(path, data_only=True, read_only=False)
    ws = wb.active
    header_row, cols = _detect_header(ws)
    headers = {c: str(ws.cell(header_row, c).value or f"col{c}") for c in range(1, ws.max_column + 1)}

    cars: list[Car] = []
    for r in range(header_row + 1, ws.max_row + 1):
        get = lambda key: ws.cell(r, cols[key]).value if key in cols else None  # noqa: E731
        price = parse_number(get("price"))
        if not price:
            continue  # пустая/итоговая строка

        brand, model = str(get("brand") or "").strip(), str(get("model") or "").strip()
        title = str(get("title") or "").strip()
        if title and not brand:
            brand, rest = split_title(title)
            model = model or rest
        elif title and not model:
            model = title

        # год: число или дата; иногда год зашит в наименование
        year_raw = get("year")
        year = getattr(year_raw, "year", None) or parse_number(year_raw)
        if not year and title:
            m = re.search(r"\b(19[89]\d|20[0-4]\d)\b", title)
            year = float(m.group()) if m else None

        mileage = parse_number(get("mileage"))
        vin = str(get("vin") or "").strip().upper()
        if not vin:  # VIN иногда лежит в соседних колонках или в примечании
            for c in range(1, ws.max_column + 1):
                m = VIN_RE.search(str(ws.cell(r, c).value or "").upper())
                if m:
                    vin = m.group()
                    break

        extra = {headers[c]: ws.cell(r, c).value for c in headers if c not in cols.values()}
        cars.append(Car(
            row=r, brand=brand, model=model,
            year=int(year) if year else None,
            mileage=int(mileage) if mileage is not None else None,
            price=price, vin=vin,
            full_vat=parse_vat(get("vat"), s.default_full_vat) if "vat" in cols else s.default_full_vat,
            extra=extra,
        ))
    return cars


# ---------------------------------------------------------------- итоговый файл

HEAD_FILL = PatternFill("solid", fgColor="15171C")
GOOD_FILL = PatternFill("solid", fgColor="E3F4EA")
BAD_FILL = PatternFill("solid", fgColor="FBE4E4")
STAGE_NAMES = {"file": "1. Файл", "market": "2. Рынок", "autoteka": "3. Автотека", "ok": "Прошла"}


def _write_table(ws, headers: list[str], rows: list[list], widths: list[int], fill=None) -> None:
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, (int, float)) and c > 3:
                cell.number_format = "# ##0"
            if fill:
                cell.fill = fill
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def is_not_passenger(e: Evaluation) -> bool:
    return e.stage == "file" and bool(e.reasons) and e.reasons[0].startswith("Не легковой")


def _lot(c: Car) -> list[str]:
    return [c.get("код предложения", "код лота", "лот", "код"), c.get("модификация"),
            c.get("федеральный округ", "регион", "город"), c.get("состояние"),
            c.get("количество ключей", "ключ"), c.get("фото")]


def _r(x):
    return round(x) if isinstance(x, (int, float)) else x


def write_result(evals: list[Evaluation], path: str | Path, s: Settings) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    total = len(evals)
    evals = [e for e in evals if not is_not_passenger(e)]  # грузовики, прицепы, спецтехника в отчёт не идут
    market_ran = any(e.stage not in ("file", "candidate") for e in evals)

    good = sorted([e for e in evals if e.passed], key=lambda e: e.profit or 0, reverse=True)
    if market_ran:
        ws = wb.create_sheet("Берём")
        _write_table(
            ws,
            ["Стр.", "Автомобиль", "VIN", "Год", "Пробег, км", "Цена в файле", "−15 %", "Платим (нал)",
             "Медиана рынка", "Продадим за", "Прибыль", "Маржа, %", "Ликвидность /10", "Продажа, дней",
             "Владельцев", "ДТП", "Код лота", "Модификация", "Регион", "Состояние", "Ключи", "Фото",
             "Заметки", "Спрос и отзывы", "Болячки", "Источники"],
            [[e.car.row, e.car.title, e.car.vin, e.car.year, e.car.mileage, _r(e.car.price), _r(e.purchase_price),
              _r(e.cash_price), e.market.median_price if e.market else None, _r(e.expected_sale), _r(e.profit),
              round(e.profit / e.cash_price * 100, 1) if e.cash_price and e.profit is not None else None,
              e.market.liquidity if e.market else None, e.market.days_to_sell if e.market else None,
              e.history.owners if e.history else "не проверено",
              e.history.accidents if e.history else "не проверено",
              *_lot(e.car), "\n".join(e.notes),
              e.market.demand_notes if e.market else "", e.market.known_issues if e.market else "",
              "\n".join(e.market.sources[:5]) if e.market else ""] for e in good],
            [6, 28, 20, 7, 11, 13, 13, 13, 13, 13, 12, 9, 10, 10, 10, 8, 10, 22, 14, 18, 14, 30, 35, 45, 40, 45],
            GOOD_FILL,
        )

    priced = sorted([e for e in evals if e.profit is not None], key=lambda e: e.profit, reverse=True)
    if priced:
        wsp = wb.create_sheet("Прибыль по всем")
        _write_table(
            wsp,
            ["Стр.", "Автомобиль", "Пробег, км", "Платим (нал)", "Рынок: от", "Рынок: медиана", "Рынок: до",
             "Продадим за", "Подготовка", "Прибыль", "Маржа, %", "Ликвидность /10", "Продажа, дней",
             "Объявлений", "Итог", "Почему нет"],
            [[e.car.row, e.car.title, e.car.mileage, _r(e.cash_price), e.market.min_price, e.market.median_price,
              e.market.max_price, _r(e.expected_sale), s.prep_cost_rub, _r(e.profit),
              round(e.profit / e.cash_price * 100, 1) if e.cash_price else None,
              e.market.liquidity, e.market.days_to_sell, e.market.listings_found,
              "БЕРЁМ" if e.passed else "нет", "\n".join(e.reasons)] for e in priced],
            [6, 28, 11, 13, 13, 13, 13, 13, 11, 12, 9, 10, 10, 10, 8, 50],
        )
        for r, e in enumerate(priced, start=2):
            wsp.cell(r, 11).number_format = "0.0"
            wsp.cell(r, 15).fill = GOOD_FILL if e.passed else BAD_FILL

    cands = [e for e in evals if e.stage == "candidate"]  # прошли бесплатный отсев, рынок не проверялся
    if cands:
        wsc = wb.create_sheet("Кандидаты")
        _write_table(
            wsc,
            ["Стр.", "Автомобиль", "VIN", "Год", "Пробег, км", "Цена в файле", "−15 %", "Платим (нал)",
             "Код лота", "Модификация", "Регион", "Состояние", "Ключи", "Фото", "Заметки"],
            [[e.car.row, e.car.title, e.car.vin, e.car.year, e.car.mileage, _r(e.car.price), _r(e.purchase_price),
              _r(e.cash_price), *_lot(e.car), "\n".join(e.notes)]
             for e in sorted(cands, key=lambda e: (e.car.brand.lower(), e.car.model.lower()))],
            [6, 28, 20, 7, 11, 13, 13, 13, 10, 22, 14, 18, 14, 30, 35],
        )

    bad = [e for e in evals if not e.passed and e.stage != "candidate"]
    ws2 = wb.create_sheet("Отсеяно (легковые)")
    _write_table(
        ws2,
        ["Стр.", "Автомобиль", "VIN", "Год", "Пробег, км", "Цена в файле", "Платим (нал)", "Этап", "Почему",
         "Прибыль", "Код лота", "Тип ТС"],
        [[e.car.row, e.car.title, e.car.vin, e.car.year, e.car.mileage, _r(e.car.price), _r(e.cash_price),
          STAGE_NAMES.get(e.stage, e.stage), "\n".join(e.reasons), _r(e.profit), e.car.get("код"),
          e.car.get("тип тс")] for e in bad],
        [6, 28, 20, 7, 11, 13, 13, 12, 60, 12, 10, 14],
        BAD_FILL,
    )

    ws3 = wb.create_sheet("Условия")
    rows = [
        ["Скидка продавца", f"{s.seller_discount:.0%}"],
        ["Коэффициент «за наличку» (полный НДС)", s.cash_factor],
        ["Макс. пробег, км", s.max_mileage_km],
        ["Макс. км в год", s.max_km_per_year],
        ["Мин. ликвидность /10", s.min_liquidity],
        ["Торг при продаже", f"{s.sale_discount:.0%}"],
        ["Подготовка и оформление, ₽", s.prep_cost_rub],
        ["Мин. прибыль, ₽", s.min_profit_rub],
        ["Год выпуска от", s.min_year or f"не старше {s.max_age_years} лет"],
        ["Тип ТС", ", ".join(s.vehicle_types)],
        ["Статус лота", ", ".join(s.allow_statuses)],
        ["Без ключей", "не рассматриваем" if s.skip_no_keys else "рассматриваем"],
        ["Правый руль", "не рассматриваем" if s.skip_rhd else "рассматриваем"],
        ["Плохое состояние", ", ".join(s.bad_conditions)],
        ["Стоп-слова в комментариях", ", ".join(s.bad_words)],
        ["Макс. владельцев", s.max_owners],
        ["Макс. ДТП", s.max_accidents],
        ["Всего строк с ценой в файле", total],
        ["Из них легковых", len(evals)],
        ["Прошли бесплатный отсев", sum(e.stage != "file" for e in evals)],
        ["Берём", len(good)],
    ]
    _write_table(ws3, ["Параметр", "Значение"], rows, [40, 60])
    wb.save(path)
