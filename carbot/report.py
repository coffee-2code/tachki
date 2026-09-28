"""Итоговый Excel: «Сводка» + листы с данными в едином стиле."""
from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Optional

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import Settings
from .models import Car, Evaluation

# ---------------------------------------------------------------- палитра и шрифты
INK = "1B1F27"
MUTED = "6B7280"
ACCENT = "F2B705"
GREEN, GREEN_BG = "1E7A4C", "E6F4EC"
RED, RED_BG = "B42318", "FDECEC"
BLUE = "1F5FBF"
ZEBRA = "F6F7F9"
LINE = "E1E4EA"
TILE_BG = "F3F4F6"
FONT = "Calibri"

HEADER_FILL = PatternFill("solid", fgColor=INK)
ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA)
BOTTOM = Border(bottom=Side(style="thin", color=LINE))

MONEY = '#,##0" ₽"'
KM = '#,##0" км"'


def f(size: int = 11, bold: bool = False, color: str = INK, **kw) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, **kw)


def plural(n: int, one: str, few: str, many: str) -> str:
    """plural(33, "машина", "машины", "машин") → «33 машины»."""
    k = n % 100
    w = many if 11 <= k <= 14 else one if n % 10 == 1 else few if 2 <= n % 10 <= 4 else many
    return f"{n} {w}"


def _money_short(x: float) -> str:
    if abs(x) >= 1_000_000:
        return f"{x / 1_000_000:.1f} млн ₽".replace(".", ",")
    return f"{x:,.0f} ₽".replace(",", " ")


# ---------------------------------------------------------------- описание колонок
@dataclass
class Col:
    title: str
    width: float
    get: Callable[[Evaluation], Any]
    kind: str = "text"  # text | wrap | int | year | money | km | pct | liq | link | badge


def _lot(c: Car, *keys: str) -> str:
    return c.get(*keys)


def _link(e: Evaluation) -> str:
    return e.market.sources[0] if e.market and e.market.sources else ""


def _margin(e: Evaluation) -> Optional[float]:
    return e.profit / e.cash_price if e.profit is not None and e.cash_price else None


def _history(e: Evaluation, attr: str) -> Any:
    return getattr(e.history, attr) if e.history else "—"


C_CAR = Col("Автомобиль", 30, lambda e: f"{e.car.brand} {e.car.model}")
C_YEAR = Col("Год", 7, lambda e: e.car.year, "year")
C_KM = Col("Пробег", 12, lambda e: e.car.mileage or None, "km")
C_REGION = Col("Регион", 16, lambda e: _lot(e.car, "федеральный округ", "регион", "город"))
C_SRS = Col("Стартовая цена\n(СРС)", 15, lambda e: e.car.price, "money")
C_MINUS = Col("После скидки\n−15 %", 15, lambda e: e.purchase_price, "money")
C_PAY = Col("Платим:\n× 0,86 (нал)", 15, lambda e: e.cash_price, "money")
C_MEDIAN = Col("Рынок: самая\nнизкая цена", 15, lambda e: e.market.median_price if e.market else None, "money")
C_SALE = Col("Продадим за\n(−5 % торг)", 15, lambda e: e.expected_sale, "money")
C_PROFIT = Col("Прибыль", 14, lambda e: e.profit, "money")
C_MARGIN = Col("Маржа", 8, _margin, "pct")
C_LIQ = Col("Ликвид-\nность", 9, lambda e: e.market.liquidity if e.market else None, "liq")
C_DAYS = Col("Продажа,\nдней", 9, lambda e: e.market.days_to_sell if e.market else None, "int")
C_ADS = Col("Объяв-\nлений", 9, lambda e: e.market.listings_found if e.market else None, "int")
def _site(e: Evaluation, name: str) -> Optional[dict]:
    for b in getattr(e.market, "by_source", None) or []:
        if b["name"] == name:
            return b
    return None


def _site_price(name: str) -> Col:
    def get(e: Evaluation):
        b = _site(e, name)
        return b["low"] if b and b.get("low") else (b["status"] if b else None)
    return Col(f"{name}:\nмин. цена", 14, get, "money")


def _site_link(name: str) -> Col:
    """Ссылка на самое дешёвое подходящее объявление, иначе — на поиск."""
    def get(e: Evaluation):
        b = _site(e, name) or {}
        return b.get("low_url") or b.get("url") or None
    return Col(f"{name}:\nобъявление ↗", 12, get, "link")


def _site_skipped(e: Evaluation) -> str:
    """Какие дешёвые объявления отбросили и почему — чтобы было видно, что проверка честная."""
    lines = []
    for b in getattr(e.market, "by_source", None) or []:
        for sk in b.get("skipped", [])[:4]:
            lines.append(f"{b['name']}: {sk}")
    return "\n".join(lines)


SITES = ("Дром", "Авто.ру")
C_SITE_PRICES = [_site_price(n) for n in SITES]
C_SITE_LINKS = [_site_link(n) for n in SITES]
C_SKIPPED = Col("Отброшены дешевле\n(почему)", 44, _site_skipped, "wrap")
C_DROM = Col("Дром", 9, _link, "link")
C_PHOTO = Col("Фото", 9, lambda e: _lot(e.car, "фото"), "link")
C_LOT = Col("Код лота", 10, lambda e: _lot(e.car, "код предложения", "код лота", "лот"))
C_VIN = Col("VIN", 20, lambda e: e.car.vin)
C_MOD = Col("Модификация", 24, lambda e: _lot(e.car, "модификация"), "wrap")
C_COND = Col("Состояние", 18, lambda e: _lot(e.car, "состояние").split("(")[0].strip())
C_KEYS = Col("Ключи", 14, lambda e: _lot(e.car, "количество ключей", "ключ"))
C_NOTES = Col("Заметки", 36, lambda e: "\n".join(e.notes), "wrap")


def _table(ws: Worksheet, title: str, subtitle: str, cols: list[Col], evals: list[Evaluation],
           min_profit: Optional[float] = None) -> None:
    """Заголовок листа (строки 1–2), шапка (3), данные с 4-й строки."""
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = f(16, True)
    ws["A2"] = subtitle
    ws["A2"].font = f(10, color=MUTED)
    ws.row_dimensions[1].height = 26

    cols = [Col("№", 5, lambda e: None, "int")] + cols
    for c, col in enumerate(cols, start=1):
        cell = ws.cell(3, c, col.title)
        cell.font = f(10, True, "FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = col.width
    ws.row_dimensions[3].height = 44

    for i, e in enumerate(evals, start=1):
        r = i + 3
        for c, col in enumerate(cols, start=1):
            v = i if c == 1 else col.get(e)
            if v in ("", None):
                v = None
            cell = ws.cell(r, c, v)
            cell.font = f(10)
            cell.border = BOTTOM
            cell.alignment = Alignment(vertical="center", wrap_text=col.kind == "wrap",
                                       indent=1 if col.kind in ("text", "wrap") and c > 1 else 0)
            if i % 2 == 0:
                cell.fill = ZEBRA_FILL
            if v is None:
                continue
            if col.kind == "money" and isinstance(v, str):
                cell.font = f(9, color=MUTED)  # «мало объявлений», «не проверено: …»
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            elif col.kind == "money":
                cell.number_format = MONEY
            elif col.kind == "km":
                cell.number_format = KM
            elif col.kind == "pct":
                cell.number_format = "0%"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col.kind in ("int", "year", "liq"):
                cell.number_format = "0"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col.kind == "link" and str(v).startswith("http"):
                cell.hyperlink = str(v)
                cell.value = "открыть ↗"
                cell.font = f(10, color=BLUE, underline="single")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col.kind == "badge":
                ok = v == "БЕРЁМ"
                cell.font = f(10, True, GREEN if ok else RED)
                cell.fill = PatternFill("solid", fgColor=GREEN_BG if ok else RED_BG)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if col.title == "Прибыль" and isinstance(v, (int, float)):
                good = min_profit is not None and v >= min_profit
                cell.font = f(10, True, GREEN if good else RED if v < 0 else INK)

    last = len(evals) + 3
    if evals:
        for c, col in enumerate(cols, start=1):
            rng = f"{get_column_letter(c)}4:{get_column_letter(c)}{last}"
            if col.kind == "liq":
                ws.conditional_formatting.add(rng, ColorScaleRule(
                    start_type="num", start_value=1, start_color="F8C9C4",
                    mid_type="num", mid_value=5, mid_color="FCE8B2",
                    end_type="num", end_value=10, end_color="B7E1C7"))
            if col.title == "Прибыль":
                ws.conditional_formatting.add(rng, DataBarRule(
                    start_type="num", start_value=0, end_type="max", color="9BD3B0", showValue=True))
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{max(last, 3)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "3:3"


# ---------------------------------------------------------------- сводка
def _tile(ws: Worksheet, col: int, label: str, value: str, color: str = INK) -> None:
    """Плитка на 2 колонки: подпись в строке 4, значение в строке 5."""
    a, b = get_column_letter(col), get_column_letter(col + 1)
    ws.merge_cells(f"{a}4:{b}4")
    ws.merge_cells(f"{a}5:{b}5")
    ws[f"{a}4"] = label
    ws[f"{a}4"].font = f(9, color=MUTED)
    ws[f"{a}5"] = value
    ws[f"{a}5"].font = f(20, True, color)
    for r in (4, 5):
        for c in (col, col + 1):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=TILE_BG)
            ws.cell(r, c).alignment = Alignment(horizontal="left", vertical="center", indent=1)


def _section(ws: Worksheet, row: int, title: str) -> int:
    ws.cell(row, 1, title).font = f(13, True)
    ws.cell(row, 1).border = Border(bottom=Side(style="medium", color=ACCENT))
    for c in range(2, 13):
        ws.cell(row, c).border = Border(bottom=Side(style="medium", color=ACCENT))
    return row + 1


def _mini_header(ws: Worksheet, row: int, titles: list[str]) -> None:
    for c, t in enumerate(titles, start=1):
        cell = ws.cell(row, c, t)
        cell.font = f(9, True, MUTED)
        cell.border = BOTTOM
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _summary(ws: Worksheet, evals: list[Evaluation], total_rows: int, source_name: str, s: Settings,
             market_ran: bool, good: list[Evaluation], priced: list[Evaluation]) -> None:
    ws.sheet_view.showGridLines = False
    for c, w in enumerate([5, 30, 8, 12, 14, 14, 14, 9, 11, 18, 12, 11], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    ws["A1"] = "ТАЧКИ · отбор машин на перепродажу"
    ws["A1"].font = f(20, True)
    ws["A2"] = (f"{source_name} · {datetime.now():%d.%m.%Y %H:%M} · цены: "
                f"{'Claude + веб-поиск' if s.market_source == 'claude' else ', '.join({'drom': 'Дром', 'autoru': 'Авто.ру'}.get(k, k) for k in s.market_sources)}")
    ws["A2"].font = f(10, color=MUTED)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[5].height = 34

    passed_file = [e for e in evals if e.stage != "file"]
    _tile(ws, 1, "Строк в файле", f"{total_rows:,}".replace(",", " "))
    _tile(ws, 3, "Легковых", str(len(evals)))
    _tile(ws, 5, "Прошли отсев", str(len(passed_file)))
    _tile(ws, 7, "Оценено на рынке", str(len(priced)) if market_ran else "—")
    _tile(ws, 9, "Берём", str(len(good)) if market_ran else "—", GREEN if good else INK)
    _tile(ws, 11, "Прибыль по «берём»", _money_short(sum(e.profit for e in good)) if good else "—",
          GREEN if good else INK)

    row = _section(ws, 7, "Лучшие по прибыли")
    if good:
        _mini_header(ws, row, ["№", "Автомобиль", "Год", "Пробег", "Платим (нал)", "Продадим за", "Прибыль",
                               "Маржа", "Ликвидность", "Регион", "Объявления", "Код лота"])
        for i, e in enumerate(good[:10], start=1):
            r = row + i
            vals = [i, f"{e.car.brand} {e.car.model}", e.car.year, e.car.mileage or None, e.cash_price,
                    e.expected_sale, e.profit, _margin(e), e.market.liquidity if e.market else None,
                    e.car.get("федеральный округ", "регион"), _link(e), e.car.get("код предложения", "код лота")]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(r, c, v)
                cell.font = f(10, bold=c == 7, color=GREEN if c == 7 else INK)
                cell.border = BOTTOM
                cell.number_format = {4: KM, 5: MONEY, 6: MONEY, 7: MONEY, 8: "0%"}.get(c, "General")
                if c in (1, 3, 8, 9):
                    cell.alignment = Alignment(horizontal="center")
                elif c in (2, 10):
                    cell.alignment = Alignment(indent=1)
                if c == 11 and v:
                    cell.hyperlink, cell.value = v, "открыть ↗"
                    cell.font = f(10, color=BLUE, underline="single")
        row += min(len(good), 10) + 2
    else:
        msg = ("Рынок ещё не проверялся — список прошедших бесплатный отсев на листе «Кандидаты»."
               if not market_ran else "Ни одна машина не прошла все условия. Что не дотянуло — на листе «Прибыль по всем».")
        ws.cell(row, 1, msg).font = f(10, color=MUTED)
        row += 2

    row = _section(ws, row, "Воронка отбора")
    _mini_header(ws, row, ["", "Этап", "Машин", "Отсеяно"])
    stages = [("Легковые в файле", len(evals))]
    stages.append(("Прошли бесплатный отсев", len(passed_file)))
    if market_ran:
        after_market = [e for e in evals if e.stage in ("autoteka", "ok")]
        stages.append(("Прошли рынок", len(after_market)))
        stages.append(("Итог: берём", len(good)))
    prev = None
    for i, (name, n) in enumerate(stages, start=1):
        r = row + i
        ws.cell(r, 2, name).font = f(10)
        ws.cell(r, 3, n).font = f(10, True)
        ws.cell(r, 4, (prev - n) if prev is not None else None).font = f(10, color=MUTED)
        for c in (2, 3, 4):
            ws.cell(r, c).border = BOTTOM
        prev = n
    first = row + 1
    ws.conditional_formatting.add(f"C{first}:C{row + len(stages)}", DataBarRule(
        start_type="num", start_value=0, end_type="max", color="F7D774", showValue=True))
    row += len(stages) + 2

    row = _section(ws, row, "Почему отсеяли")
    reasons = Counter(e.reasons[0].split(":")[0] for e in evals if e.reasons and not e.passed
                      and e.stage != "candidate")
    _mini_header(ws, row, ["", "Причина", "Машин"])
    for i, (name, n) in enumerate(reasons.most_common(), start=1):
        r = row + i
        ws.cell(r, 2, name).font = f(10)
        ws.cell(r, 3, n).font = f(10, True)
        ws.cell(r, 2).border = ws.cell(r, 3).border = BOTTOM
    if reasons:
        ws.conditional_formatting.add(f"C{row + 1}:C{row + len(reasons)}", DataBarRule(
            start_type="num", start_value=0, end_type="max", color="F4A9A0", showValue=True))
    row += len(reasons) + 2

    row = _section(ws, row, "Условия отбора")
    conds = [
        ("Цена закупки", f"стартовая цена (СРС) − {s.seller_discount:.0%}, затем × {s.cash_factor} "
                         "(наличка, полный НДС)"),
        ("Рынок", f"Дром и Авто.ру, машины до {str(s.max_distance_km)[:-3] + ' ' + str(s.max_distance_km)[-3:]} км от Москвы (примерно по дорогам): самые "
                  "дешёвые объявления того же года с похожим пробегом (±35 %); каждое открыто и прочитано — учёт "
                  "только РФ, в наличии (не под заказ / в пути), не битая. Рынок — самая низкая подходящая цена. "
                  "Вывод о выгоде — только если проверены обе площадки"
         if s.market_source != "claude" else "Claude с веб-поиском по площадкам и форумам"),
        ("Прибыль", f"рынок × {1 - s.sale_discount:.2f} (торг) − цена закупки − {s.prep_cost_rub:,} ₽ подготовка"
                    .replace(",", " ")),
        ("Лоты", f"тип: {', '.join(s.vehicle_types)}; статус: {', '.join(s.allow_statuses)}"
                 + ("; без правого руля" if s.skip_rhd else "") + ("; с ключами" if s.skip_no_keys else "")
                 + ("; без VIN США/Канады/Мексики" if s.skip_us_vin else "")),
        ("Состояние", "не " + ", ".join(s.bad_conditions) + "; в комментариях нет: " + ", ".join(s.bad_words)),
        ("Год и пробег", f"от {s.min_year or s.current_year - s.max_age_years} г.; от {s.min_mileage_km} "
                         f"до {s.max_mileage_km:,} км (0–1 км — не заводится); "
                         f"до {s.max_km_per_year:,} км/год".replace(",", " ")),
        ("Пороги", f"ликвидность от {s.min_liquidity}/10; прибыль от {s.min_profit_rub:,} ₽".replace(",", " ")),
        ("Автотека", f"владельцев до {s.max_owners}; ДТП до {s.max_accidents}; без скрутки"),
    ]
    for i, (k, v) in enumerate(conds, start=1):
        r = row + i - 1
        ws.cell(r, 2, k).font = f(10, True)
        ws.cell(r, 2).alignment = Alignment(vertical="top")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=12)
        ws.cell(r, 3, v).font = f(10)
        ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * max(1, -(-len(v) // 105)) + 2  # по числу строк текста
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


# ---------------------------------------------------------------- сборка
def is_not_passenger(e: Evaluation) -> bool:
    return e.stage == "file" and bool(e.reasons) and e.reasons[0].startswith("Не легковой")


STAGE_NAMES = {"file": "Файл", "market": "Рынок", "autoteka": "Автотека", "ok": "Прошла"}


def write_result(evals: list[Evaluation], path: str | Path, s: Settings, source_name: str = "") -> None:
    total = len(evals)
    evals = [e for e in evals if not is_not_passenger(e)]  # грузовики, прицепы, спецтехника в отчёт не идут
    market_ran = any(e.stage not in ("file", "candidate") for e in evals)
    good = sorted([e for e in evals if e.passed], key=lambda e: e.profit or 0, reverse=True)
    priced = sorted([e for e in evals if e.profit is not None], key=lambda e: e.profit, reverse=True)
    cands = sorted([e for e in evals if e.stage == "candidate"],
                   key=lambda e: (e.car.brand.lower(), e.car.model.lower(), e.car.year or 0))
    bad = sorted([e for e in evals if not e.passed and e.stage not in ("candidate",)],
                 key=lambda e: (e.reasons[0].split(":")[0] if e.reasons else "", e.car.brand.lower()))

    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.sheet_properties.tabColor = ACCENT
    _summary(ws, evals, total, source_name or Path(path).stem, s, market_ran, good, priced)

    if market_ran:
        w = wb.create_sheet("Берём")
        w.sheet_properties.tabColor = GREEN
        profit = _money_short(sum(e.profit for e in good)) if good else "0 ₽"
        _table(w, f"Берём — {plural(len(good), 'машина', 'машины', 'машин')}, прибыль {profit}",
               "Прошли все этапы. Отсортировано по прибыли. VIN проверьте в Автотеке перед покупкой, "
               "если она не подключена.",
               [C_CAR, C_YEAR, C_KM, C_REGION, C_SRS, C_MINUS, C_PAY, *C_SITE_PRICES, C_MEDIAN, C_SALE,
                C_PROFIT, C_MARGIN, C_LIQ, C_DAYS, *C_SITE_LINKS, C_PHOTO, Col("Владель-\nцев", 9, lambda e: _history(e, "owners"), "int"),
                Col("ДТП", 6, lambda e: _history(e, "accidents"), "int"),
                C_LOT, C_VIN, C_MOD, C_COND, C_KEYS, C_NOTES],
               good, s.min_profit_rub)

    if priced:
        w = wb.create_sheet("Прибыль по всем")
        w.sheet_properties.tabColor = BLUE
        _table(w, f"Прибыль по всем оценённым — {len(priced)}",
               f"Порог прибыли {s.min_profit_rub:,} ₽, ликвидности {s.min_liquidity}/10. ".replace(",", " ")
               + "Здесь видно и то, что не дотянуло.",
               [Col("Итог", 9, lambda e: "БЕРЁМ" if e.passed else "нет", "badge"),
                C_CAR, C_YEAR, C_KM, C_SRS, C_MINUS, C_PAY, *C_SITE_PRICES, C_MEDIAN, C_SALE, C_PROFIT, C_MARGIN,
                C_LIQ, C_ADS, *C_SITE_LINKS,
                Col("Почему нет", 50, lambda e: "\n".join(e.reasons), "wrap"), C_SKIPPED],
               priced, s.min_profit_rub)

    if cands:
        w = wb.create_sheet("Кандидаты")
        w.sheet_properties.tabColor = "9CA3AF"
        _table(w, f"Кандидаты — {len(cands)}",
               "Прошли бесплатный отсев, рынок по ним не проверялся.",
               [C_CAR, C_YEAR, C_KM, C_REGION, C_SRS, C_MINUS, C_PAY, C_COND, C_KEYS, C_PHOTO, C_LOT, C_VIN, C_MOD,
                C_NOTES],
               cands)

    w = wb.create_sheet("Отсеяно")
    w.sheet_properties.tabColor = RED
    _table(w, f"Отсеяно — {plural(len(bad), 'легковая', 'легковые', 'легковых')}",
           "Грузовики, прицепы и спецтехника в отчёт не включены. Фильтр по колонке «Причина» — стрелка в шапке.",
           [C_CAR, Col("Причина", 20, lambda e: e.reasons[0].split(":")[0] if e.reasons else ""),
            Col("Подробно", 44, lambda e: "\n".join(e.reasons), "wrap"),
            C_YEAR, C_KM, C_SRS, C_MINUS, C_PAY, C_PROFIT,
            Col("Этап", 9, lambda e: STAGE_NAMES.get(e.stage, e.stage)), C_LOT, C_VIN],
           bad)

    wb.save(path)
