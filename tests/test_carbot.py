import asyncio
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from carbot import pricing
from carbot.autoteka import detect_rollback
from carbot.config import Settings
from carbot.excel_io import parse_vat, read_cars, split_title, write_result
from carbot.filters import file_stage_reasons
from carbot.models import Car, HistoryReport, MarketReport
from carbot.pipeline import evaluate


@pytest.fixture
def s():
    st = Settings()
    st.autoteka_client_id = st.autoteka_client_secret = ""
    st.blacklist = []
    st.min_profit_rub = 150_000
    return st


def car(**kw):
    base = dict(row=2, brand="Kia", model="K5", year=Settings().current_year - 3, mileage=45_000,
                price=3_000_000, vin="XWEG2417BN0012345", full_vat=True)
    base.update(kw)
    return Car(**base)


def test_prices(s):
    c = car(price=1_000_000)
    assert pricing.purchase_price(c, s) == pytest.approx(850_000)
    assert pricing.cash_price(c, s) == pytest.approx(731_000)          # 850 000 × 0,86
    assert pricing.cash_price(car(price=1_000_000, full_vat=False), s) == pytest.approx(850_000)


def test_file_filters(s):
    assert file_stage_reasons(car(), s) == []
    assert "Старая" in file_stage_reasons(car(year=2010), s)[0]
    assert "Большой пробег" in file_stage_reasons(car(mileage=260_000), s)[0]
    assert "км/год" in file_stage_reasons(car(year=s.current_year - 2, mileage=70_000), s)[0]
    s.blacklist = ["lada"]
    assert file_stage_reasons(car(brand="Lada", model="Granta"), s)


def test_rollback():
    assert not detect_rollback([("2021-01-01", 10_000), ("2022-01-01", 30_000)], 3000)
    assert detect_rollback([("2021-01-01", 90_000), ("2023-05-01", 60_000)], 3000)
    assert not detect_rollback([("2021-01-01", 50_000), ("2021-02-01", 48_500)], 3000)  # погрешность


def test_parse_helpers():
    assert parse_vat("С НДС 20%", False) is True
    assert parse_vat("без НДС", True) is False
    assert split_title("Land Rover Defender 110") == ("Land Rover", "Defender 110")
    assert split_title("Toyota Camry 2.5") == ("Toyota", "Camry 2.5")


def test_read_messy_excel(tmp_path, s):
    wb = Workbook()
    ws = wb.active
    ws.append(["Реестр ТС на реализацию"])  # шапка-заголовок над таблицей
    ws.append([])
    ws.append(["№", "Наименование ТС", "Год выпуска", "Пробег, км", "VIN", "Стоимость, руб. с НДС", "НДС"])
    ws.append([1, "Toyota Camry", 2021, "68 000", "XW7BF4FK30S123456", "3 150 000 ₽", "полный"])
    ws.append([2, "Land Rover Discovery Sport", 2019, 90000, "SALCA2BN5KH812345", 2500000, "без НДС"])
    ws.append([None, "Итого", None, None, None, None, None])
    p = tmp_path / "in.xlsx"
    wb.save(p)

    cars = read_cars(p, s)
    assert len(cars) == 2
    a, b = cars
    assert (a.brand, a.model, a.year, a.mileage, a.price, a.full_vat) == ("Toyota", "Camry", 2021, 68000, 3150000, True)
    assert a.vin == "XW7BF4FK30S123456"
    assert (b.brand, b.model, b.full_vat) == ("Land Rover", "Discovery Sport", False)


class FakeMarket:
    def __init__(self, reports):
        self.reports = reports

    async def analyze(self, c):
        return self.reports[c.row]


class FakeAutoteka:
    enabled = True

    def __init__(self, hist):
        self.hist = hist

    async def check(self, vin):
        return self.hist[vin]


def mr(median, liq=8, n=20):
    return MarketReport(median_price=median, min_price=median - 200_000, max_price=median + 200_000,
                        listings_found=n, liquidity=liq, days_to_sell=20, demand_notes="спрос есть",
                        known_issues="", sources=["https://auto.ru/x"])


def test_pipeline(tmp_path, s):
    y = s.current_year
    cars = [
        car(row=2, vin="AAAAAAAAAAAAAAAA1", price=2_000_000),                  # хорошая
        car(row=3, vin="AAAAAAAAAAAAAAAA2", year=2009),                        # старая
        car(row=4, vin="AAAAAAAAAAAAAAAA3", price=2_000_000),                  # неликвид
        car(row=5, vin="AAAAAAAAAAAAAAAA4", price=2_000_000),                  # мало прибыли
        car(row=6, vin="AAAAAAAAAAAAAAAA5", price=2_000_000),                  # 3 владельца
        car(row=7, vin="AAAAAAAAAAAAAAAA6", price=2_000_000, mileage=40_000),  # скрутка
    ]
    # платим 2 000 000 × 0,85 × 0,86 = 1 462 000
    market = FakeMarket({2: mr(2_000_000), 4: mr(2_000_000, liq=3), 5: mr(1_600_000),
                         6: mr(2_000_000), 7: mr(2_000_000)})
    hist = {
        "AAAAAAAAAAAAAAAA1": HistoryReport(owners=1, accidents=0, last_known_mileage=45_000),
        "AAAAAAAAAAAAAAAA5": HistoryReport(owners=3, accidents=0),
        "AAAAAAAAAAAAAAAA6": HistoryReport(owners=1, accidents=0, mileage_rollback=True, last_known_mileage=120_000),
    }
    for c in cars:
        c.model = f"K{c.row}"  # разные модели — разные запросы к рынку
    cars.append(car(row=8, model="K2", vin="AAAAAAAAAAAAAAAA7", price=2_000_000))  # та же модель, что row 2
    hist["AAAAAAAAAAAAAAAA7"] = HistoryReport(owners=1, accidents=2)
    calls = []
    orig = market.analyze

    async def counting(c):
        calls.append(c.market_key)
        return await orig(c)
    market.analyze = counting
    evals = asyncio.run(evaluate(cars, s, market, FakeAutoteka(hist)))
    assert len(calls) == 5  # row 8 взяла оценку row 2 из кэша
    by = {e.car.row: e for e in evals}
    assert by[2].passed and by[2].profit == pytest.approx(2_000_000 * 0.95 - 1_462_000 - s.prep_cost_rub)
    assert by[3].stage == "file" and not by[3].passed
    assert "Неликвид" in by[4].reasons[0]
    assert "Мало прибыли" in by[5].reasons[0]
    assert "Владельцев" in by[6].reasons[0]
    assert any("Скрученный" in r for r in by[7].reasons)
    assert "ДТП" in by[8].reasons[0] and by[8].market is by[2].market
    assert [e.car.row for e in evals if e.passed] == [2]

    out = tmp_path / "out.xlsx"
    write_result(evals, out, s)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Берём", "Прибыль по всем", "Отсеяно (легковые)", "Условия"]
    wp = wb["Прибыль по всем"]
    assert wp.max_row == 1 + 6  # все, кто прошёл рынок, включая отсеянных потом
    profits = [wp.cell(r, 10).value for r in range(2, wp.max_row + 1)]
    assert profits == sorted(profits, reverse=True)
    assert wp.cell(2, 15).value in ("БЕРЁМ", "нет")
    assert wb["Берём"].max_row == 2 and wb["Отсеяно (легковые)"].max_row == 7


def test_leasing_registry_filters(s):
    ok = car(extra={"Тип ТС": "ЛЕГКОВОЙ", "Статус ИЗТ": "В продаже", "Состояние ПЛ": "Хорошее (сколы)",
                    "Комментарий по оценке": "Хорошее состояние ТС. 2 ключа."})
    assert file_stage_reasons(ok, s) == []
    truck = car(extra={"Тип ТС": "ГРУЗОВОЙ"})
    assert file_stage_reasons(truck, s)[0].startswith("Не легковой")
    reserved = car(extra={"Тип ТС": "ЛЕГКОВОЙ", "Статус ИЗТ": "Резерв"})
    assert "Не в продаже" in file_stage_reasons(reserved, s)[0]
    assert "Не в продаже" in file_stage_reasons(car(extra={"Статус ИЗТ": "Оценка"}), s)[0]
    assert file_stage_reasons(car(extra={"Статус ИЗТ": "В продаже"}), s) == []
    rhd = car(model="LAND CRUISER PRADO (правый руль)", extra={"Статус ИЗТ": "В продаже"})
    assert file_stage_reasons(rhd, s) == ["Правый руль"]
    assert file_stage_reasons(car(extra={"Нет ключей": "Нет ключей"}), s) == ["Нет ключей"]
    assert file_stage_reasons(car(extra={"Количество ключей после изъятия": "Нет ключей"}), s) == ["Нет ключей"]
    assert file_stage_reasons(car(extra={"Нет ключей": "", "Количество ключей после изъятия": "1 ключ"}), s) == []
    hard = car(extra={"Тип ТС": "ЛЕГКОВОЙ", "Состояние ПЛ": "HARD (тотал, сгоревшие)"})
    assert "Состояние" in file_stage_reasons(hard, s)[0]
    crashed = car(extra={"Тип ТС": "ЛЕГКОВОЙ", "Комментарии по ключам": "На ходу, после ДТП, замята крыша"})
    assert "после дтп" in file_stage_reasons(crashed, s)[0]
    assert "Старая" in file_stage_reasons(car(year=2017), s)[0]


def test_registry_columns(tmp_path, s):
    wb = Workbook()
    ws = wb.active
    ws.append([]); ws.append([])
    ws.append(["Код предложения", "Фото и видео материалы ТС", "Марка", "Модель", "Модификация", "VIN", "Тип ТС",
               "Пробег", "Год выпуска", "СРС с переоценкой", "Комментарий ОРИТ"])
    ws.append(["014782", "https://disk", "Chery", "Tiggo 8 Pro Max", "Ultimate 4WD", "LVTDD24B9PD144841",
               "ЛЕГКОВОЙ", 74308, 2023, 2_500_000, ""])
    p = tmp_path / "reg.xlsx"
    wb.save(p)
    (c,) = read_cars(p, s)
    assert (c.brand, c.model, c.year, c.mileage, c.price) == ("Chery", "Tiggo 8 Pro Max", 2023, 74308, 2_500_000)
    assert c.get("код предложения") == "014782" and c.get("тип тс") == "ЛЕГКОВОЙ"
    assert c.market_key == ("chery", "tiggo 8 pro max", 2023, "ultimate 4wd")


def test_report_only_passenger_cars(tmp_path, s):
    from carbot.pipeline import stage_file
    cars = [car(row=2, extra={"Тип ТС": "ЛЕГКОВОЙ"}), car(row=3, extra={"Тип ТС": "ГРУЗОВОЙ"}),
            car(row=4, year=2015, extra={"Тип ТС": "ЛЕГКОВОЙ"})]
    out = tmp_path / "c.xlsx"
    write_result(stage_file(cars, s), out, s)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Кандидаты", "Отсеяно (легковые)", "Условия"]  # рынок не запускали — «Берём» нет
    assert wb["Кандидаты"].max_row == 2 and wb["Отсеяно (легковые)"].max_row == 2  # грузовика нет нигде
